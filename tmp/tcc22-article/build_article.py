from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


REFERENCE = Path(r"C:\Users\Pichau\Downloads\Modelo_artigo_Computacao_e_software_2026_2-6a8d952b18e6b (1).docx")
OUTPUT = Path(r"C:\Users\Pichau\Documents\TCC\PrograMago\docs\PrograMago_artigo_introducao_metodos.docx")


def font(style, size=12, bold=None):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is not None:
        rfonts.set(qn("w:ascii"), "Arial")
        rfonts.set(qn("w:hAnsi"), "Arial")
        rfonts.set(qn("w:eastAsia"), "Arial")


doc = Document(REFERENCE)
body = doc._body._body
for element in list(body):
    if element.tag != qn("w:sectPr"):
        body.remove(element)

section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(3)
section.bottom_margin = Cm(2)
section.left_margin = Cm(3)
section.right_margin = Cm(2)

for name in ("Normal", "Title", "Assinaturas", "Heading 1", "Heading 2", "1 Texto", "4 Referências", "Título não numerado"):
    if name in doc.styles:
        font(doc.styles[name])

font(doc.styles["Title"], bold=True)
font(doc.styles["Heading 1"], bold=True)
font(doc.styles["Heading 2"], bold=False)
font(doc.styles["1 Texto"], bold=False)
font(doc.styles["4 Referências"], bold=False)

body_style = doc.styles["1 Texto"]
body_style.paragraph_format.first_line_indent = Cm(2)
body_style.paragraph_format.line_spacing = 1.5
body_style.paragraph_format.space_before = Pt(0)
body_style.paragraph_format.space_after = Pt(0)

for name in ("Heading 1", "Heading 2"):
    s = doc.styles[name]
    s.paragraph_format.first_line_indent = Cm(0)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.space_before = Pt(12)
    s.paragraph_format.space_after = Pt(12)
    s.paragraph_format.line_spacing = 1.5

refs_style = doc.styles["4 Referências"]
refs_style.paragraph_format.first_line_indent = Cm(0)
refs_style.paragraph_format.line_spacing = 1.0
refs_style.paragraph_format.space_after = Pt(12)


def add_body(text):
    p = doc.add_paragraph(style="1 Texto")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(text)
    return p


def add_heading(text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(text)
    return p


title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(18)
title.add_run(
    "DESENVOLVIMENTO E AVALIAÇÃO DE UM PROTÓTIPO DE JOGO EDUCACIONAL "
    "PARA A APRENDIZAGEM DE CLASSES, OBJETOS, ENCAPSULAMENTO, HERANÇA E POLIMORFISMO"
)

author = doc.add_paragraph(style="Assinaturas")
author.alignment = WD_ALIGN_PARAGRAPH.RIGHT
author.paragraph_format.space_after = Pt(18)
run = author.add_run("João Vitor Rodrigues Rocha")
run.bold = True
run.font.name = "Arial"
run.font.size = Pt(12)

add_heading("INTRODUÇÃO")

for paragraph in [
    "Na disciplina introdutória de programação orientada a objetos (POO) da Universidade Federal da Paraíba, campus IV, os índices de reprovação superaram 50% na maioria dos oito semestres examinados por Henrique e Rebouças (2015). As autoras também registraram abandono superior a 24% em todos os períodos analisados, entre 2009.2 e 2013.1. Trata-se de um resultado localizado, que não representa todas as disciplinas de POO, mas evidencia a necessidade de investigar as dificuldades de aprendizagem nesse conteúdo.",
    "A POO organiza programas por meio de classes e objetos, articulando estado e comportamento. Para utilizá-la, o estudante precisa distinguir a definição de uma classe das instâncias criadas a partir dela e compreender como atributos, métodos, encapsulamento, herança e polimorfismo se relacionam durante a execução. Em levantamento com 74 estudantes que haviam cursado POO, Henrique e Rebouças (2015) receberam 32 indicações de polimorfismo entre os conteúdos considerados difíceis. Ao examinar as respostas de estudantes a atividades de orientação a objetos, Abbasi et al. (2021) identificaram confusões entre classes, objetos, atributos e métodos, além de dificuldades na construção de hierarquias de herança.",
    "A dificuldade aparece em outros contextos, embora com intensidades distintas. Em uma disciplina de POO com Java da Universidad Nacional Tecnológica de Lima Sur, no Peru, Díaz-Leyva e Chamorro-Atalaya (2020) registraram reprovação de 64% e 40% em dois períodos de 2019. Na pesquisa com 25 estudantes de um desses períodos, a aprendizagem dos conceitos foi a dimensão percebida como mais difícil. O curso analisado não possuía carga horária teórica, circunstância que impede atribuir os resultados apenas à natureza da POO. Em estudo brasileiro mais recente, Araújo et al. (2025) observaram baixo desempenho em testes aplicados a 162 estudantes de um curso técnico que incluía disciplinas de lógica e de POO; esse resultado descreve o conjunto investigado, não uma taxa isolada de reprovação em POO.",
    "Propostas pedagógicas têm buscado tornar os conceitos observáveis e oferecer oportunidades de prática. Figueiredo et al. (2015) relataram a aplicação de uma abordagem gamificada em uma disciplina que ensinava orientação a objetos. Em estudo experimental com 83 estudantes, Abbasi et al. (2021) compararam o uso de um jogo sério a uma forma convencional de ensino e observaram maior ganho de aprendizagem no grupo que utilizou o jogo. Esses resultados justificam investigar recursos lúdicos, mas não demonstram antecipadamente que qualquer novo jogo produzirá o mesmo efeito.",
    "Nesse contexto, o PrograMago é proposto como um protótipo de jogo educacional em que o estudante escreve construções de código semelhantes às de Java e observa, em uma arena, seus efeitos sobre personagens e batalhas. O jogo será organizado em desafios progressivos sobre classes, objetos, encapsulamento, herança e polimorfismo. Ao relacionar o código produzido às mudanças de estado e comportamento dos objetos representados visualmente, a proposta busca oferecer uma forma adicional de praticar conceitos que exigem abstração. Sua contribuição será examinada por meio da análise do artefato e da experiência dos participantes, sem pressupor eficácia antes da avaliação.",
    "A pesquisa parte da seguinte pergunta: de que maneira a utilização de um protótipo de jogo educacional pode apoiar a compreensão e a prática dos conceitos de classes, objetos, encapsulamento, herança e polimorfismo por estudantes de programação? O objetivo geral consiste em desenvolver e avaliar um protótipo que relacione a escrita de código a efeitos visuais em uma arena. Para alcançá-lo, serão examinadas dificuldades descritas na literatura, projetadas e implementadas as atividades do jogo, verificada a correção do conteúdo com especialistas e analisados indícios de aprendizagem, motivação e usabilidade em uma aplicação com estudantes.",
    "O estudo tem natureza aplicada e caráter exploratório e descritivo, com abordagem quantitativa e qualitativa. O protótipo será desenvolvido e submetido a testes técnicos e a uma avaliação de conteúdo antes da aplicação com estudantes. Nessa aplicação, serão utilizados um teste anterior, um teste posterior, registros de interação e um questionário de percepção. A combinação desses dados permitirá descrever o desempenho e as dificuldades observadas no grupo participante, bem como identificar ajustes necessários no jogo."
]:
    add_body(paragraph)

add_heading("PROCEDIMENTO EXPERIMENTAL")
add_body(
    "O procedimento compreende a construção do protótipo, sua verificação técnica, a validação do conteúdo e uma aplicação exploratória com estudantes. As etapas serão executadas nessa ordem para que problemas de funcionamento ou de correção conceitual sejam tratados antes da coleta de dados com o público-alvo."
)

add_heading("MATERIAIS", 2)
for paragraph in [
    "O principal material da pesquisa será o PrograMago, protótipo de jogo educacional 2D desenvolvido na Unity 6. A interface reunirá um editor de código, orientações da tarefa, mensagens de erro e uma arena visual. O jogador escreverá um subconjunto de construções com aparência de Java, limitado aos conceitos previstos no estudo. O sistema analisará esse texto por regras próprias e atualizará a arena quando a resposta for válida; o código digitado pelo participante não será compilado nem executado como um programa Java completo.",
    "O conteúdo será distribuído em quatro fases. A primeira abordará classe, objeto, atributos, construtor e encapsulamento por meio da criação de um mago. A segunda acrescentará métodos e interação com objetos inimigos. A terceira trabalhará herança e sobrescrita com subclasses de magos. A quarta apresentará polimorfismo pela atribuição de objetos de subclasses a uma referência do tipo da classe-base e pela observação de comportamentos distintos na batalha. Cada fase disponibilizará explicações, objetivos, validação da resposta e dicas progressivas.",
    "Serão utilizados um roteiro para avaliação por professores ou profissionais de programação, dois testes com questões equivalentes sobre os conceitos abordados e um questionário de percepção dos estudantes. O roteiro examinará correção conceitual, clareza das instruções, adequação dos exemplos de código e coerência entre código e resposta visual. Os testes verificarão identificação e aplicação de classes, objetos, atributos, métodos, encapsulamento, herança e polimorfismo. O questionário reunirá itens fechados e respostas abertas sobre facilidade de uso, clareza, motivação, feedback e dificuldades encontradas. Quando tecnicamente disponíveis, os registros do jogo incluirão fase concluída, tentativas, erros, dicas utilizadas e tempo de interação.",
    "A população de interesse compreende estudantes com contato prévio com programação e com os conteúdos introdutórios de POO. A amostra será não probabilística, por conveniência e adesão voluntária, composta pelos estudantes que aceitarem participar da aplicação. Não se fixará previamente um tamanho amostral, pois a investigação tem finalidade exploratória; a quantidade efetiva de participantes e suas características serão informadas na apresentação dos resultados."
]:
    add_body(paragraph)

add_heading("MÉTODOS", 2)
for paragraph in [
    "Quanto à natureza, a pesquisa será aplicada, pois produzirá e examinará um artefato destinado ao apoio à aprendizagem. Quanto aos objetivos, será exploratória e descritiva. A avaliação com estudantes adotará um delineamento de grupo único com pré-teste e pós-teste, acompanhado de questionário e registros de uso. Não haverá grupo de controle nesta etapa; por isso, diferenças entre os testes serão tratadas como indícios observados no grupo, e não como prova causal da eficácia do jogo.",
    "A primeira etapa consistirá em revisar estudos empíricos sobre dificuldades de aprendizagem de POO e intervenções com jogos educacionais. Com base no escopo definido para o protótipo, serão elaborados os desafios, as explicações, as regras de análise do código e as respostas visuais da arena. A implementação na Unity será seguida de testes com entradas válidas e inválidas, cobrindo declarações de classes, instanciação, atributos privados, construtores, métodos, herança, sobrescrita e polimorfismo. Também serão verificados os retornos apresentados ao estudante e a progressão entre as fases.",
    "Após a verificação técnica, professores ou profissionais da área avaliarão o protótipo com o roteiro de conteúdo. As observações serão registradas por tópico e usadas para corrigir erros conceituais ou instruções ambíguas. Uma aplicação piloto verificará se as tarefas, os instrumentos e o fluxo de uso podem ser concluídos sem bloqueios relevantes. Os ajustes necessários serão realizados antes da aplicação principal.",
    "Na aplicação com estudantes, inicialmente serão apresentados os objetivos da pesquisa e registradas informações sobre contato prévio com programação e POO. Depois do consentimento, cada participante responderá ao pré-teste, utilizará o protótipo na sequência prevista de fases, responderá ao pós-teste e preencherá o questionário de percepção. As questões dos dois testes cobrirão os mesmos conceitos com enunciados equivalentes, sem repetir literalmente as respostas esperadas. O tempo de uso e os eventos registrados pelo jogo serão vinculados a um identificador de participante, sem exposição do nome nos dados analisados.",
    "As respostas dos testes serão corrigidas por uma chave previamente definida e convertidas em escores comparáveis. Serão descritas as frequências de acerto por conceito e a diferença entre os escores anterior e posterior de cada participante; medidas-resumo serão apresentadas de acordo com o número de casos obtido. As respostas fechadas do questionário e os registros do jogo serão sintetizados por frequências. As respostas abertas e os comentários dos avaliadores serão agrupados por temas, como compreensão dos conceitos, clareza do feedback e dificuldades de uso. A interpretação cruzará esses dados para identificar convergências e divergências, respeitando o caráter exploratório e a ausência de grupo de controle.",
    "Antes de recrutar estudantes, serão confirmadas com o orientador e com a instituição as exigências éticas aplicáveis à pesquisa com participantes. A participação será voluntária, dependerá de consentimento informado e poderá ser interrompida sem prejuízo ao estudante. Os dados utilizados na análise serão identificados por código e acessíveis apenas aos responsáveis pela pesquisa."
]:
    add_body(paragraph)

refs_heading = doc.add_paragraph(style="Título não numerado")
refs_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
refs_heading.paragraph_format.keep_with_next = True
refs_heading.paragraph_format.space_before = Pt(18)
refs_heading.paragraph_format.space_after = Pt(12)
refs_run = refs_heading.add_run("REFERÊNCIAS")
refs_run.bold = True
refs_run.font.name = "Arial"
refs_run.font.size = Pt(12)

references = [
    "ABBASI, Suhni; KAZI, Hameedullah; KAZI, Ahmed Waliullah; KHOWAJA, Kamran; BALOCH, Ahsanullah. Gauge object oriented programming in student’s learning performance, normalized learning gains and perceived motivation with serious games. Information, Basel, v. 12, n. 3, art. 101, 2021. DOI: https://doi.org/10.3390/info12030101. Acesso em: 13 set. 2026.",
    "ARAÚJO, Ana Lívia dos Santos; OLIVEIRA, João Gabriel de Araújo; SOUZA, Ilane Alícia Fernandes de; AZEVEDO, Giovanna Silva de; BATISTA, Cybelly Francisca Silva; PEREIRA, Rafael Peixoto de Moraes. Aprendizagem e retenção em programação: um estudo no ensino médio técnico do Instituto Federal de Educação, Ciência e Tecnologia do Rio Grande do Norte, campus Parelhas, Brasil. Revista Principia, João Pessoa, v. 62, e8545, 2025. DOI: https://doi.org/10.18265/2447-9187a2024id8545. Acesso em: 13 set. 2026.",
    "DÍAZ-LEYVA, Teodoro; CHAMORRO-ATALAYA, Omar. Analysis of learning difficulties in object oriented programming in systems engineering students at UNTELS. Advances in Science, Technology and Engineering Systems Journal, v. 5, n. 6, p. 1704-1709, 2020. DOI: https://doi.org/10.25046/aj0506203. Acesso em: 13 set. 2026.",
    "FIGUEIREDO, Karen da Silva; RIBEIRO, Jivago Medeiros; SOUZA, Raphael; ANGELO, Vinicius Raniero. Uma abordagem gamificada para o ensino de programação orientada a objetos. In: WORKSHOP SOBRE EDUCAÇÃO EM COMPUTAÇÃO, 23., 2015, Recife. Anais [...]. Porto Alegre: Sociedade Brasileira de Computação, 2015. p. 316-325. DOI: https://doi.org/10.5753/wei.2015.10248. Acesso em: 13 set. 2026.",
    "HENRIQUE, Mychelline Souto; REBOUÇAS, Ayla Débora Dantas Souza. Objetos de aprendizagem para auxiliar o ensino de conceitos do paradigma de programação orientada a objetos. RENOTE, Porto Alegre, v. 13, n. 2, 2015. DOI: https://doi.org/10.22456/1679-1916.61433. Acesso em: 13 set. 2026."
]
for reference in references:
    prefix, remainder = reference.split("DOI: ", 1)
    doi_url, access = remainder.split(". Acesso em:", 1)
    reference = (
        f"{prefix}DOI: {doi_url.removeprefix('https://doi.org/')}. "
        f"Disponível em: {doi_url}. Acesso em:{access}"
    )
    p = doc.add_paragraph(style="4 Referências")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(reference)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
